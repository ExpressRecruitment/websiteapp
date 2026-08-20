#!/usr/bin/env python3
"""
Regenerate data/jobs.json from an "ER Web" export sheet.

The source is a spreadsheet with a sheet named "ER Web" containing one
column, four rows per job, scraped straight from the website's jobs
listing page in this order: Title, "Location <x>", "Job Type <x>",
"Salary <x>". That's the same shape as the 2026-08-14 export this
snapshot was generated from.

Usage:
    python3 scripts/generate-jobs-data.py path/to/export.xlsx [snapshot-date]

    snapshot-date defaults to today (YYYY-MM-DD) if not given.

Requires: openpyxl (`pip install openpyxl`)
"""
import json
import re
import sys
import unicodedata
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
OUT_PATH = REPO_ROOT / "data" / "jobs.json"

# Express Recruitment's real division list (from the source workbook's
# "Current Jobs" sheet, division dropdown column).
DIVISIONS = [
    "Temporary",
    "Sales",
    "Professional & Corporate Support",
    "Legal & Finance",
    "Not-For-Profit",
    "Tech & IT",
    "Engineering",
    "Executive & Managerial Search",
    "The Academy",
]


def clean(s):
    if s is None:
        return ""
    s = s.replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip()


def slugify(title, seen):
    base = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    base = re.sub(r"[^a-zA-Z0-9]+", "-", base).strip("-").lower()
    base = re.sub(r"-{2,}", "-", base) or "role"
    slug = base
    n = 2
    while slug in seen:
        slug = f"{base}-{n}"
        n += 1
    seen.add(slug)
    return slug


def classify_division(title):
    """Best-effort division guess from the job title alone - the source
    export doesn't carry a real division per job. Check against the real
    taxonomy (Express Recruitment > Formidable/WordPress) if/when it's
    available instead of relying on this."""
    t = title.lower()
    if t.startswith("temporary") or t.startswith("temp "):
        return "Temporary"
    if re.search(r"\b(apprentice|graduate|trainee)\b|no experience needed", t):
        return "The Academy"
    if re.search(r"\b(director|head of|managing director|general manager)\b", t):
        return "Executive & Managerial Search"
    if re.search(r"charity|not[- ]for[- ]profit|\bngo\b|fundraising", t):
        return "Not-For-Profit"
    if re.search(
        r"\b(engineer|cnc|electrician|technician|fitter|welder|mechanic|"
        r"maintenance|toolmaker|fabricat\w*|machinist|miller|programmer/operator)\b",
        t,
    ):
        return "Engineering"
    if re.search(
        r"\b(it|software|developer|sql|database|network|helpdesk|service desk|"
        r"cyber|system administrator|systems analyst)\b",
        t,
    ):
        return "Tech & IT"
    if re.search(
        r"account(ant|s|ing)?\b|finance|financial|payroll|legal|solicitor|"
        r"paralegal|credit control|bookkeep\w*|audit",
        t,
    ):
        return "Legal & Finance"
    if re.search(r"\bsales\b|business development|telesales|account executive|account manager\b", t):
        return "Sales"
    return "Professional & Corporate Support"


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = Path(sys.argv[1])
    snapshot_date = sys.argv[2] if len(sys.argv) > 2 else date.today().isoformat()

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["ER Web"]
    rows = [r[0] for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True)]

    jobs = []
    seen_slugs = set()
    for i in range(0, len(rows) - 3, 4):
        title, loc_raw, jt_raw, sal_raw = rows[i : i + 4]
        title = clean(title)
        if not title:
            continue
        location = clean(loc_raw).replace("Location", "", 1).strip()
        job_type = clean(jt_raw).replace("Job Type", "", 1).strip()
        salary = clean(sal_raw).replace("Salary", "", 1).strip()
        division = classify_division(title)
        slug = slugify(title, seen_slugs)

        jobs.append(
            {
                "id": len(jobs) + 1,
                "slug": slug,
                "title": title,
                "excerpt": f"{location} - {salary}" if salary else location,
                "description": (
                    f"{title}\n\n"
                    f"Location: {location}\n"
                    f"Job type: {job_type}\n"
                    f"Salary: {salary}\n\n"
                    "For the full role description and to apply, see this vacancy on the "
                    "Express Recruitment website."
                ),
                "link": "https://www.express-recruitment.co.uk/jobs/",
                "date": snapshot_date,
                "divisions": [division],
                "location": location or None,
                "salary": salary or None,
                "jobType": job_type or None,
            }
        )

    print(f"Parsed {len(jobs)} jobs")
    print(Counter(j["divisions"][0] for j in jobs))

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT_PATH, "w") as f:
        json.dump(jobs, f, indent=2, ensure_ascii=False)
    print("Wrote", OUT_PATH)


if __name__ == "__main__":
    main()
